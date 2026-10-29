"""TELELAUDO T6 — importação assistida de modelos DOCX.

Converte o documento em seções do modelo (técnica, achados, conclusão) e devolve um **relatório de
incompatibilidades**. O resultado é sempre um **rascunho**: publicar continua exigindo a conferência do texto e da
pré-visualização (T3). Nada do documento é executado.

Segurança: só DOCX (macro ``.docm`` recusado), limite de tamanho, de entradas e de tamanho descompactado (bomba de
compressão), XML sem resolução de entidades (python-docx), sem objetos incorporados nem conteúdo ativo.

Formatação (``rich=True``, padrão): negrito, itálico, sublinhado, quebras de linha, listas com marcadores e numeradas,
títulos e tabelas (com células mescladas na horizontal) são preservados em texto rico seguro (``richtext``). Cores,
fontes, tamanhos, imagens e objetos não são importados (listados no relatório): a aparência vem da identidade visual.
Cabeçalho e rodapé do documento são extraídos como sugestão para a identidade visual (T2), não aplicados.
``rich=False`` mantém a conversão anterior em texto simples.
"""

from __future__ import annotations

import io
import re
import zipfile
from typing import Any

from vyonix.errors import ValidationFailed

MAX_BYTES = 5 * 1024 * 1024
MAX_UNZIPPED = 50 * 1024 * 1024
MAX_ENTRIES = 2000
_HEAD = {"tecnica": "technique", "técnica": "technique", "achados": "findings", "relatorio": "findings",
         "relatório": "findings", "descricao": "findings", "descrição": "findings", "analise": "findings",
         "análise": "findings", "conclusao": "conclusion", "conclusão": "conclusion", "impressao": "conclusion",
         "impressão": "conclusion", "opiniao": "conclusion", "opinião": "conclusion"}
_MERGE = re.compile(r"MERGEFIELD\s+([^\s\\]+)", re.I)
_MAP_MERGE = {"paciente": "paciente.nome", "nome": "paciente.nome", "nomepaciente": "paciente.nome",
              "nascimento": "paciente.nascimento", "datanascimento": "paciente.nascimento", "sexo": "paciente.sexo",
              "idade": "paciente.idade", "modalidade": "exame.modalidade", "exame": "exame.descricao",
              "dataexame": "exame.data", "data": "exame.data", "accession": "exame.accession",
              "solicitante": "exame.solicitante", "medicosolicitante": "exame.solicitante",
              "clinica": "instituicao.nome", "instituicao": "instituicao.nome", "unidade": "unidade.nome"}


def _check_package(data: bytes) -> zipfile.ZipFile:
    if not data or len(data) > MAX_BYTES:
        raise ValidationFailed("Arquivo vazio ou acima de 5 MB.")
    if data[:2] != b"PK":
        raise ValidationFailed("O arquivo não é um DOCX (Word 2007 ou mais recente).")
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
        infos = z.infolist()
    except zipfile.BadZipFile:
        raise ValidationFailed("DOCX danificado.") from None
    if len(infos) > MAX_ENTRIES or sum(i.file_size for i in infos) > MAX_UNZIPPED:
        raise ValidationFailed("DOCX recusado: conteúdo compactado grande demais.")
    names = {i.filename for i in infos}
    if "word/document.xml" not in names:
        raise ValidationFailed("O arquivo não é um documento do Word (DOCX).")
    if any(n.lower().endswith("vbaproject.bin") for n in names):
        raise ValidationFailed("Documento com macros (DOCM) não é aceito. Salve como DOCX sem macros.")
    return z


def _heading(text: str) -> str | None:
    """Título de seção: o parágrafo inteiro (“Achados”) ou o trecho antes de dois-pontos (“Conclusão: …”)."""
    t = text.strip().lower()
    head = t.split(":", 1)[0] if ":" in t else t
    return _HEAD.get(re.sub(r"[\-–—.\s]+$", "", head).strip())


def _on(el, tag: str) -> bool:
    """Propriedade booleana do Word (``<w:b/>``, ``<w:b w:val="0"/>``…)."""
    from docx.oxml.ns import qn
    x = el.find(qn(tag)) if el is not None else None
    if x is None:
        return False
    v = x.get(qn("w:val"))
    return v is None or v.lower() not in ("0", "false", "none", "off")


def _runs_html(p_el) -> str:
    """Parágrafo do Word → HTML inline (negrito/itálico/sublinhado/quebras), texto escapado."""
    import html as _html
    from docx.oxml.ns import qn
    out = []
    for r in p_el.iter(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        b, i, u = _on(rpr, "w:b"), _on(rpr, "w:i"), _on(rpr, "w:u")
        buf = []
        for x in r:
            if x.tag == qn("w:t"):
                buf.append(_html.escape(x.text or "", quote=False))
            elif x.tag == qn("w:tab"):
                buf.append(" ")
            elif x.tag in (qn("w:br"), qn("w:cr")):
                buf.append("<br>")
        t = "".join(buf)
        if not t:
            continue
        if u:
            t = f"<u>{t}</u>"
        if i:
            t = f"<em>{t}</em>"
        if b:
            t = f"<strong>{t}</strong>"
        out.append(t)
    h = "".join(out)
    for tag in ("strong", "em", "u"):                      # junta trechos vizinhos com a mesma formatação
        h = h.replace(f"</{tag}><{tag}>", "")
    return h


def _list_kind(p, numbering) -> str | None:
    """``ul``/``ol`` se o parágrafo é item de lista (numeração do Word ou estilo de lista), senão ``None``."""
    from docx.oxml.ns import qn
    num = p._p.find(qn("w:pPr") + "/" + qn("w:numPr"))
    style = (p.style.name if p.style is not None else "") or ""
    if num is None and "List" not in style:
        return None
    if "Number" in style:
        return "ol"
    if "Bullet" in style:
        return "ul"
    if num is not None and numbering is not None:
        try:
            num_id = num.find(qn("w:numId")).get(qn("w:val"))
            ilvl = (num.find(qn("w:ilvl")).get(qn("w:val")) if num.find(qn("w:ilvl")) is not None else "0")
            n = next(x for x in numbering.iter(qn("w:num")) if x.get(qn("w:numId")) == num_id)
            aid = n.find(qn("w:abstractNumId")).get(qn("w:val"))
            a = next(x for x in numbering.iter(qn("w:abstractNum")) if x.get(qn("w:abstractNumId")) == aid)
            lvl = next(x for x in a.iter(qn("w:lvl")) if x.get(qn("w:ilvl")) == ilvl)
            fmt = lvl.find(qn("w:numFmt")).get(qn("w:val"))
            return "ul" if fmt in ("bullet", "none") else "ol"
        except (StopIteration, AttributeError):
            return "ul"
    return "ul"


def _heading_level(p) -> str | None:
    style = ((p.style.name if p.style is not None else "") or "").lower()
    if style.startswith(("heading", "título", "titulo")) or style == "title":
        m = re.search(r"(\d+)", style)
        return "h3" if not m or int(m.group(1)) <= 2 else "h4"
    return None


def convert(data: bytes, *, rich: bool = True) -> dict[str, Any]:
    """DOCX → {technique, findings, conclusion, format, header, footer, issues[], fields[], name}. Com ``rich`` as
    seções vêm em HTML seguro (``format="html"``); sem, em texto simples (``format="text"``)."""
    import docx
    from docx.oxml.ns import qn
    z = _check_package(data)
    names = [i.filename for i in z.infolist()]
    issues: list[str] = []
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ValidationFailed(f"Não foi possível ler o DOCX ({type(exc).__name__}).") from None
    sections = {"technique": [], "findings": [], "conclusion": []}
    rich_sec = {"technique": [], "findings": [], "conclusion": []}
    lists: dict[str, list] = {k: [] for k in rich_sec}         # lista aberta por seção: [kind, [itens]]
    current, fields, formatting = "findings", set(), False
    try:
        numbering = d.part.numbering_part.element
    except (NotImplementedError, KeyError, AttributeError):
        numbering = None

    def close_list(sec: str) -> None:
        if lists[sec]:
            kind, items = lists[sec]
            rich_sec[sec].append(f"<{kind}>" + "".join(f"<li>{x}</li>" for x in items) + f"</{kind}>")
            lists[sec] = []

    def merge_fix(text: str, html_text: str, child) -> tuple[str, str]:
        for instr in child.iter(qn("w:instrText")):
            m = _MERGE.search(instr.text or "")
            if m:
                key = re.sub(r"[^a-z]", "", m.group(1).lower())
                if key in _MAP_MERGE:
                    fields.add(_MAP_MERGE[key])
                    ph = "{{" + _MAP_MERGE[key] + "}}"
                    text = text.replace(f"«{m.group(1)}»", ph)
                    html_text = html_text.replace(f"«{m.group(1)}»", ph)
                    if ph not in text:
                        text += " " + ph
                    if ph not in html_text:
                        html_text += " " + ph
                else:
                    issues.append(f"Campo de mesclagem “{m.group(1)}” sem correspondência: virou texto.")
        return text, html_text
    body = d.element.body
    for child in body.iterchildren():
        tag = child.tag
        if tag == qn("w:p"):
            from docx.text.paragraph import Paragraph
            p = Paragraph(child, d)
            text, inline = merge_fix(p.text, _runs_html(child), child)
            if any(r.bold or r.italic or r.underline or (r.font.color and r.font.color.rgb) for r in p.runs):
                formatting = True
            if any(r.font.color and r.font.color.rgb for r in p.runs):
                issues.append("Cores do texto não são importadas (a aparência vem da identidade visual).")
            if child.find(".//" + qn("w:drawing")) is not None or child.find(".//" + qn("w:pict")) is not None:
                issues.append("Imagem no corpo do documento foi ignorada (use a logomarca/papel timbrado da identidade).")
            if child.find(".//" + qn("w:object")) is not None:
                issues.append("Objeto incorporado foi ignorado (não é executado nem importado).")
            h = _heading(text) if len(text.split(":", 1)[0]) < 40 else None
            if h:
                close_list(current)
                current = h
                rest = text.split(":", 1)[1].strip() if ":" in text else ""
                if rest:
                    sections[current].append(rest)
                    import html as _html
                    rich_sec[current].append(f"<p>{_html.escape(rest, quote=False)}</p>")
                continue
            sections[current].append(text)
            kind = _list_kind(p, numbering)
            if kind:
                if lists[current] and lists[current][0] != kind:
                    close_list(current)
                if not lists[current]:
                    lists[current] = [kind, []]
                lists[current][1].append(inline)
                continue
            close_list(current)
            hl = _heading_level(p)
            rich_sec[current].append(f"<{hl}>{inline}</{hl}>" if hl else f"<p>{inline}</p>")
        elif tag == qn("w:tbl"):
            from docx.table import Table
            close_list(current)
            t = Table(child, d)
            rows_html = []
            for tr in child.iter(qn("w:tr")):
                cells = []
                for tc in tr.iter(qn("w:tc")):
                    span = tc.find(qn("w:tcPr") + "/" + qn("w:gridSpan"))
                    n = int(span.get(qn("w:val"))) if span is not None and (span.get(qn("w:val")) or "").isdigit() else 1
                    paras = [_runs_html(x) for x in tc.iter(qn("w:p"))]
                    attr = f' colspan="{min(n, 10)}"' if n > 1 else ""
                    cells.append(f"<td{attr}>" + "<br>".join(x for x in paras if x) + "</td>")
                rows_html.append("<tr>" + "".join(cells) + "</tr>")
            rich_sec[current].append("<table><tbody>" + "".join(rows_html) + "</tbody></table>")
            for row in t.rows:
                sections[current].append(" | ".join(c.text.strip() for c in row.cells))
            if not rich:
                issues.append("Tabela convertida em linhas de texto separadas por “|”.")
        elif tag == qn("w:sdt"):
            close_list(current)
            txt = "".join(x.text or "" for x in child.iter(qn("w:t")))
            sections[current].append(txt)
            import html as _html
            rich_sec[current].append(f"<p>{_html.escape(txt, quote=False)}</p>")
            issues.append("Controle de conteúdo (formulário) convertido em texto simples.")
    for k in lists:
        close_list(k)
    if formatting and not rich:
        issues.append("Negrito, itálico, sublinhado e cores não são preservados nesta versão (o modelo é texto por "
                      "seções).")
    if any(n.startswith("word/embeddings/") for n in names):
        issues.append("O documento contém arquivos incorporados: foram ignorados.")
    fonts = sorted({r.font.name for p in d.paragraphs for r in p.runs if r.font.name})
    if fonts:
        issues.append(f"Fontes do documento ({', '.join(fonts[:6])}) não são aplicadas: a fonte vem da identidade "
                      "visual.")
    header = " ".join(p.text.strip() for s in d.sections for p in s.header.paragraphs if p.text.strip())
    footer = " ".join(p.text.strip() for s in d.sections for p in s.footer.paragraphs if p.text.strip())
    if header or footer:
        issues.append("Cabeçalho e rodapé do documento foram extraídos como sugestão para a identidade visual "
                      "(não aplicados ao modelo).")

    def clean(lines: list[str]) -> str:
        text = "\n".join(lines)
        return re.sub(r"\n{3,}", "\n\n", text).strip()
    out = {k: clean(v) for k, v in sections.items()}
    if not any(out.values()):
        raise ValidationFailed("O documento não tem texto aproveitável.")
    if rich:
        from vyonix.report import richtext

        def clean_html(parts: list[str]) -> str:
            h = "".join(parts)
            h = re.sub(r"^(?:<p></p>)+|(?:<p></p>)+$", "", h)           # parágrafos vazios nas bordas
            h = re.sub(r"(?:<p></p>){2,}", "<p></p>", h)
            return richtext.sanitize(h)
        out = {k: clean_html(v) for k, v in rich_sec.items()}
    title = (d.core_properties.title or "").strip()
    return {**out, "format": "html" if rich else "text", "header": header[:300], "footer": footer[:300], "fields": sorted(fields),
            "issues": list(dict.fromkeys(issues)), "name": title[:120]}

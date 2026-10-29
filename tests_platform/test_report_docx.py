"""TELELAUDO T6 — importação assistida de DOCX: seções reconhecidas, incompatibilidades relatadas, rascunho que
exige conferência para publicar, e recusa de macros, bombas de compressão e arquivos que não são DOCX.
Documentos sintéticos gerados no teste."""

from __future__ import annotations

import io
import unittest
import zipfile

from tests_platform.support import require_db
from tests_platform.test_pacs import ApiBase


def setUpModule():
    global db
    db = require_db()
    from tests_platform import test_pacs
    test_pacs.db = db


def make_docx() -> bytes:
    import docx
    d = docx.Document()
    d.core_properties.title = "TC de abdome (importado)"
    d.sections[0].header.paragraphs[0].text = "Clínica Exemplo — Cabeçalho"
    d.sections[0].footer.paragraphs[0].text = "Rua Exemplo, 1 — Rodapé"
    d.add_paragraph("Técnica:")
    d.add_paragraph("Aquisição com contraste em {{paciente.nome}}.")
    p = d.add_paragraph("Achados")
    p.runs[0].bold = True
    d.add_paragraph("Fígado de dimensões normais. Ver {{paciente.cpf}}.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Órgão", "Medida"
    t.cell(1, 0).text, t.cell(1, 1).text = "Baço", "10 cm"
    d.add_paragraph("Conclusão: exame dentro da normalidade.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


class DocxImport(unittest.TestCase):
    def test_convert_sections_and_issues(self):
        """Modo texto simples (compatibilidade): tabelas e formatação viram texto e são relatadas."""
        from vyonix.report import docx_import
        r = docx_import.convert(make_docx(), rich=False)
        self.assertEqual(r["format"], "text")
        self.assertEqual(r["technique"], "Aquisição com contraste em {{paciente.nome}}.")
        self.assertIn("Fígado de dimensões normais.", r["findings"])
        self.assertIn("Baço | 10 cm", r["findings"])
        self.assertEqual(r["conclusion"], "exame dentro da normalidade.")
        self.assertEqual((r["header"], r["footer"], r["name"]), ("Clínica Exemplo — Cabeçalho", "Rua Exemplo, 1 — Rodapé",
                                                                 "TC de abdome (importado)"))
        joined = " ".join(r["issues"])
        for want in ("Tabela convertida", "Negrito", "Cabeçalho e rodapé"):
            self.assertIn(want, joined)

    def test_convert_preserves_word_formatting(self):
        """Modo padrão: negrito, itálico, sublinhado, listas, títulos e tabelas (com mesclagem) viram HTML seguro."""
        import docx
        from vyonix.report import docx_import, richtext
        d = docx.Document()
        d.add_paragraph("Técnica:")
        p = d.add_paragraph()
        p.add_run("Negrito").bold = True
        p.add_run(" e ")
        p.add_run("itálico").italic = True
        p.add_run(" e ")
        p.add_run("sublinhado").underline = True
        d.add_paragraph("Achados")
        d.add_paragraph("Item A", style="List Bullet")
        d.add_paragraph("Item B", style="List Bullet")
        d.add_paragraph("Passo 1", style="List Number")
        d.add_heading("Fígado", level=2)
        t = d.add_table(rows=2, cols=3)
        t.cell(0, 0).merge(t.cell(0, 1))
        t.cell(0, 0).text, t.cell(1, 2).text = "Mesclada", "a<b"
        d.add_paragraph("Conclusão: normal.")
        buf = io.BytesIO()
        d.save(buf)
        r = docx_import.convert(buf.getvalue())
        self.assertEqual(r["format"], "html")
        self.assertEqual(r["technique"], "<p><strong>Negrito</strong> e <em>itálico</em> e <u>sublinhado</u></p>")
        self.assertIn("<ul><li>Item A</li><li>Item B</li></ul><ol><li>Passo 1</li></ol><h3>Fígado</h3>", r["findings"])
        self.assertIn('<td colspan="2">Mesclada</td>', r["findings"])
        self.assertIn("a&lt;b", r["findings"])                                   # texto escapado, não marcação
        self.assertEqual(r["conclusion"], "<p>normal.</p>")
        self.assertEqual(richtext.sanitize(r["findings"]), r["findings"])        # já dentro da lista permitida
        self.assertFalse(any("Negrito" in i or "Tabela" in i for i in r["issues"]))

    def test_rejects_unsafe_or_invalid(self):
        from vyonix.errors import ValidationFailed
        from vyonix.report import docx_import

        def zipped(entries: dict) -> bytes:
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
                for n, d in entries.items():
                    z.writestr(n, d)
            return buf.getvalue()
        cases = {"macro": zipped({"word/document.xml": "<x/>", "word/vbaProject.bin": b"\0" * 10}),
                 "bomba": zipped({"word/document.xml": "<x/>", "a.bin": b"\0" * (60 * 1024 * 1024)}),
                 "não docx": b"%PDF-1.7 qualquer coisa", "zip sem documento": zipped({"x.txt": "oi"}),
                 "vazio": b""}
        for name, data in cases.items():
            with self.assertRaises(ValidationFailed, msg=name):
                docx_import.convert(data)


class DocxApi(ApiBase):
    def test_import_creates_draft_that_requires_review(self):
        c = self.login()
        r = c.post("/api/v1/report/templates/import-docx", params={"modality": "CT"}, content=make_docx(),
                   headers={"Content-Type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"})
        self.assertEqual(r.status_code, 201, r.text)
        d = r.json()
        self.assertEqual((d["status"], d["version"]), ("draft", 1))
        self.assertTrue(any("paciente.cpf" in i for i in d["issues"]))           # campo desconhecido virou texto
        items = c.get("/api/v1/report/templates", params={"include_inactive": True}).json()["items"]
        it = next(x for x in items if x["id"] == d["id"])
        self.assertIsNone(it["published_version"])
        self.assertIn("[paciente.cpf]", it["draft"]["findings"])
        self.assertEqual(it["draft"]["content_format"], "html")
        self.assertIn("<table>", it["draft"]["findings"])                       # tabela do Word preservada
        self.assertTrue(it["draft"]["import_report"]["issues"])
        self.assertEqual(c.post(f"/api/v1/report/templates/{d['id']}/publish", json={}).status_code, 422)
        self.assertEqual(c.post(f"/api/v1/report/templates/{d['id']}/publish", json={"reviewed": True}).status_code, 200)


if __name__ == "__main__":
    unittest.main()

"""Gera docs/homologacao/VYONIX_P0_Acompanhamento.xlsx a partir de P0_acompanhamento.json (fonte única,
sincronizada com VYONIX_NEXT_ACTION.md). Uso: python tools/p0_planilha.py"""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "homologacao" / "P0_acompanhamento.json"
OUT = ROOT / "docs" / "homologacao" / "VYONIX_P0_Acompanhamento.xlsx"
F = "Arial"
THIN = Side(style="thin", color="B7C4CC")
B = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="0F3D56")
EDIT = PatternFill("solid", fgColor="FFF2CC")


def header(ws, row, cols):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row=row, column=i, value=c)
        x.font, x.fill, x.border = Font(name=F, bold=True, color="FFFFFF"), HEAD, B
        x.alignment = Alignment(wrap_text=True, vertical="center")


def body(ws, r, vals, edit=()):
    for i, v in enumerate(vals, 1):
        x = ws.cell(row=r, column=i, value=v)
        x.font, x.border = Font(name=F, size=10), B
        x.alignment = Alignment(wrap_text=True, vertical="top")
        if i in edit:
            x.fill = EDIT


def widths(ws, ws_w):
    for col, w in ws_w.items():
        ws.column_dimensions[col].width = w


def main() -> Path:
    d = json.loads(SRC.read_text(encoding="utf-8"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Requisitos P0"
    ws["A1"] = f"VYONIX — P0: requisitos e situação (atualizado em {d['atualizado']})"
    ws["A1"].font = Font(name=F, bold=True, size=13)
    ws["A2"] = (f"Commit da implementação: {d['commit']} (ramo develop). Teste local ≠ CI Windows ≠ validação no "
                "computador do proprietário. Nada é marcado como aprovado sem execução. Células amarelas: atualizar.")
    ws["A2"].font = Font(name=F, size=9, italic=True)
    header(ws, 4, ["#", "Requisito", "O que mudou", "Testes locais", "CI Windows", "Validação no computador do "
                   "proprietário", "Evidência", "Data da validação"])
    for i, q in enumerate(d["requisitos"], 1):
        body(ws, 4 + i, [i, q["req"], q["mudou"], q["local"], q["ci"], q["prop"], q["evid"], q["data"]],
             edit=(4, 5, 6, 7, 8))
    last = 4 + len(d["requisitos"])
    dv = DataValidation(type="list", formula1='"Aprovado,Reprovado,Não executado,Pendente"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"D5:F{last}")
    s = last + 2
    ws.cell(row=s, column=2, value="Resumo").font = Font(name=F, bold=True)
    header(ws, s + 1, ["", "Situação", "", "Testes locais", "CI Windows", "Proprietário"])
    for k, st in enumerate(["Aprovado", "Reprovado", "Não executado", "Pendente"]):
        r = s + 2 + k
        body(ws, r, ["", st, "", f"=COUNTIF(D$5:D${last},$B{r})", f"=COUNTIF(E$5:E${last},$B{r})",
                     f"=COUNTIF(F$5:F${last},$B{r})"])
    widths(ws, dict(zip("ABCDEFGH", (4, 32, 60, 14, 15, 20, 48, 13))))
    ws.freeze_panes = "A5"

    w2 = wb.create_sheet("Execuções Windows")
    w2["A1"] = "Execuções do CI Windows"
    w2["A1"].font = Font(name=F, bold=True, size=13)
    header(w2, 3, ["Run", "Versão", "Commit", "Situação", "Resultado / observação", "Evidência", "Data"])
    for i, e in enumerate(d["execucoes"]):
        body(w2, 4 + i, [e["run"], e["versao"], e["commit"], e["situacao"], e["obs"], e["evid"], e["data"]],
             edit=(4, 5, 6, 7))
    dv2 = DataValidation(type="list", formula1='"Aprovado,Reprovado,Em execução,Pendente"')
    w2.add_data_validation(dv2)
    dv2.add("D4:D60")
    widths(w2, dict(zip("ABCDEFG", (14, 20, 10, 13, 70, 50, 12))))

    w3 = wb.create_sheet("Próximos passos")
    w3["A1"] = "Próximos passos"
    w3["A1"].font = Font(name=F, bold=True, size=13)
    header(w3, 3, ["#", "Passo", "Condição", "Feito?"])
    for i, p in enumerate(d["passos"], 1):
        body(w3, 3 + i, [i, p["passo"], p["condicao"], p["feito"]], edit=(4,))
    dv3 = DataValidation(type="list", formula1='"Sim,Não"')
    w3.add_data_validation(dv3)
    dv3.add("D4:D60")
    w3.cell(row=5 + len(d["passos"]), column=2,
            value="Nada foi alterado no computador do proprietário sem autorização.").font = Font(name=F, size=9, italic=True)
    widths(w3, dict(zip("ABCD", (4, 70, 50, 10))))
    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    print(main())
